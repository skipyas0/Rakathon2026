"""Závěr zasedání MDT: docs/template_zaver_zasedani/form.md filled for one meeting, as PDF (preview
and download) and DOCX (download, for edits).

HEADER   – the board (record) and the patient (FHIR Patient). No AI.
BODY     – the patient's history: one LLM call over the FHIR record (POST …/zaver/draft), one value
           per template line plus the resources it came from, cached in data/zavery/<case>.json.
           AI output: blue in both formats, its sources under each value. "Otázka na MDT" comes
           from the nomination (record), not the LLM.
RESPONSE – the board's confirmed decision (record). No AI.
The coordinator previews the PDF; to change anything they download the DOCX and edit it."""

import datetime as dt
import io
import json
import tempfile
from dataclasses import dataclass
from pathlib import Path
from zoneinfo import ZoneInfo

import httpx
from docx import Document
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response

from mdt_backend import agent, llm
from mdt_backend.config import settings
from mdt_backend.features import cases, records
from mdt_backend.pdf import BLUE, GREY, PdfDoc, Run
from mdt_backend.provenance import NOT_FOUND

router = APIRouter(prefix="/api/cases")

TEMPLATE = Path(__file__).parents[4] / "docs" / "template_zaver_zasedani" / "form.md"
DRAFTS = Path(__file__).parents[4] / "data" / "zavery"  # <case id>.json, AI history per meeting
PRAGUE = ZoneInfo("Europe/Prague")  # stored times are UTC; the record shows local time
NOT_FILLED = "nevyplněno"  # no AI draft for this field (yet): not a claim that data are missing
NO_SOURCE = "⚠ bez ověřeného zdroje"
QUESTION = "Otázka na MDT"

INTENTS = {
    "curative": "Kurativní",
    "palliative": "Paliativní",
    "supportive": "Podpůrná péče",
    "surveillance": "Sledování",
    "deferred": "Odloženo",
}
# template department line → the board specialty's label (as the confirmed decision's attendees
# carry it, frontend SPECIALTIES); the coordinator line is whoever confirmed the conclusion.
# ponytail: the template comes from a sarcoma board; other specialties go on one extra line.
DEPARTMENTS = {
    "Klinika zobrazovacích metod": "Radiologie",
    "Onkologická klinika": "Klinická onkologie",
    "III.chirurgická klinika": "Chirurgie",
    "Ústav Patologie": "Patologie",
}
COORDINATOR = "Koordinátorka onkologické péče"


@dataclass
class Field:
    label: str
    depth: int
    value: str  # template default ("?", "0", "") until filled
    group: bool  # has indented children
    path: str  # "Historie > Staging > CT"
    sources: list[str] | None = None  # None = not AI-filled


def parse_template(text: str) -> dict[str, list[Field | None]]:
    """`# NAME` sections → lines; None = blank line (kept for layout)."""
    sections: dict[str, list[Field | None]] = {}
    current: list[Field | None] = []
    for line in text.splitlines():
        if line.startswith("# "):
            current = sections.setdefault(line[2:].strip(), [])
        elif not line.strip():
            current.append(None)
        else:
            label, _, value = line.strip().partition(":")
            current.append(
                Field(label.strip(), len(line) - len(line.lstrip("\t")), value.strip(), False, "")
            )
    for lines in sections.values():
        stack: list[Field] = []
        fields = [f for f in lines if f]
        for i, f in enumerate(fields):
            f.group = i + 1 < len(fields) and fields[i + 1].depth > f.depth
            stack = [s for s in stack if s.depth < f.depth]
            f.path = " > ".join([s.label for s in stack] + [f.label])
            stack.append(f)
    return sections


def leaves(lines: list[Field | None]) -> list[Field]:
    return [f for f in lines if f and not f.group]


# --- HEADER: board and patient ---------------------------------------------------------------


def fill_header(
    lines: list[Field | None], patient: dict, d: records.Decision, board: str | None
) -> None:
    present = {a.specialty: a for a in d.attendees}
    patient_values = {
        "Jméno, příjmení pacienta": (patient.get("name") or [{}])[0].get("text", NOT_FOUND),
        # the dataset has no birth number, insurer, address or phone: say so, never invent them
        "Rodné číslo (pojišťovna) pacienta": NOT_FOUND,
        "Adresa pacienta": NOT_FOUND,
        "Telefonní kontakt na pacienta": NOT_FOUND,
    }
    for f in (f for f in lines if f):
        if f.label.endswith("multidisciplinární tým"):  # "Sarcomaboard, multidisciplinární tým"
            f.label = f"{board or 'Komise'}, multidisciplinární tým"
        elif f.label in patient_values:
            f.value = patient_values[f.label]
        elif f.label in DEPARTMENTS:
            a = present.get(DEPARTMENTS[f.label])
            f.value = (a.name or "přítomen") if a else "–"
        elif f.label == COORDINATOR:
            f.value = d.confirmed_by or "–"
    others = [s for s in present if s not in DEPARTMENTS.values()]
    if others:
        at = next(i for i, f in enumerate(lines) if f and f.label == COORDINATOR)
        lines.insert(at, Field("Další odbornosti", 0, ", ".join(others), False, "Další odbornosti"))


# --- BODY: the patient's history, drafted by the LLM -----------------------------------------

HISTORY_SYSTEM = f"""Jsi klinický dokumentarista onkologického multidisciplinárního týmu (MDT).
Z dodaného FHIR záznamu pacienta vyplň pole formuláře pro zasedání MDT. Pravidla:
- Jen fakta ze záznamu, nic nedomýšlej. Nejsou-li data, napiš přesně: "{NOT_FOUND}"
- Extrémně stručně: heslovité body oddělené "; ", žádné celé věty (epikríza max. 3 krátké body).
- Čísla vždy, kde to jde: data (MM/RRRR nebo DD.MM.RRRR), velikosti v mm, hodnoty s jednotkami,
  počty cyklů, délky v měsících, TNM, ECOG, procenta.
- Pole s výchozí hodnotou 0 jsou počty/přítomnost: "0" = nebylo; jinak počet nebo stručný popis
  (např. "1; 03/2025 jaterní meta"). "Délka" dispenzarizace v měsících.
- "zdroje": ID FHIR zdrojů (ResourceType/id přesně jak jsou v záznamu), ze kterých hodnota pochází.
  U "{NOT_FOUND}" prázdný seznam.
- Data obsahují záměrné chyby a mezery: rozpory uveď, neopravuj je.
- Česky, standardní lékařské zkratky."""


def _history_schema(fields: list[Field]) -> dict:
    item = {
        "type": "object",
        "properties": {
            "hodnota": {"type": "string"},
            "zdroje": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["hodnota", "zdroje"],
        "additionalProperties": False,
    }
    props = {f"f{i:02d}": {**item, "description": f.path} for i, f in enumerate(fields)}
    return {
        "type": "object",
        "properties": props,
        "required": list(props),
        "additionalProperties": False,
    }


def draft_history(resources: list[dict]) -> dict[str, dict]:
    """Template path → {value, sources} for every BODY line but the question; one LLM call.
    Sources the record doesn't contain are dropped."""
    fields = [
        f
        for f in leaves(parse_template(TEMPLATE.read_text("utf-8"))["BODY"])
        if f.label != QUESTION
    ]
    field_list = "\n".join(
        f"f{i:02d}: {f.path}" + (f" (výchozí: {f.value})" if f.value else "")
        for i, f in enumerate(fields)
    )
    record = "\n".join(
        json.dumps(
            {k: v for k, v in r.items() if k != "meta"}, ensure_ascii=False, separators=(",", ":")
        )
        for r in resources
    )
    user = f"Pole formuláře:\n{field_list}\n\nFHIR záznam pacienta (1 zdroj na řádek):\n{record}"
    answer = llm.chat_json(HISTORY_SYSTEM, user, _history_schema(fields), name="historie")
    known = {f"{r['resourceType']}/{r['id']}" for r in resources}
    return {
        f.path: {
            "value": answer[f"f{i:02d}"]["hodnota"].strip() or NOT_FOUND,
            "sources": [s for s in answer[f"f{i:02d}"]["zdroje"] if s in known],
        }
        for i, f in enumerate(fields)
    }


def _draft_path(case_id: str) -> Path:
    return DRAFTS / f"{case_id}.json"


def load_draft(case_id: str) -> dict | None:
    path = _draft_path(case_id)
    return json.loads(path.read_text("utf-8")) if path.exists() else None


@router.post("/{case_id}/zaver/draft")
def make_draft(case_id: str, refresh: bool = False) -> dict:
    """The AI's draft of the history fields: cached, made on first request (~30 s)."""
    patient = cases.meetings().get(case_id)
    if patient is None:
        raise HTTPException(404, f"No case {case_id!r}")
    if not refresh and (draft := load_draft(case_id)):
        return {"created": draft["created"]}
    if not settings.azure_api_key:
        raise HTTPException(503, "AI návrh nelze vytvořit: chybí přístup k jazykovému modelu")
    with cases.fhir_server():
        resources = list(agent.Record(patient).bundle.by_ref.values())  # dataset ids, as cited
    try:
        fields = draft_history(resources)
    except httpx.HTTPStatusError as e:  # e.g. 401 = the LLM key is invalid or expired
        raise HTTPException(
            503, f"AI návrh nelze vytvořit (model: HTTP {e.response.status_code})"
        ) from e
    except (httpx.HTTPError, RuntimeError, KeyError) as e:
        raise HTTPException(503, f"AI návrh nelze vytvořit ({type(e).__name__})") from e
    draft = {
        "created": dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
        "model": settings.llm_model,
        "fields": fields,
    }
    DRAFTS.mkdir(parents=True, exist_ok=True)
    _draft_path(case_id).write_text(json.dumps(draft, ensure_ascii=False, indent=2), "utf-8")
    return {"created": draft["created"]}


# --- the filled template -------------------------------------------------------------------


def _local(t: dt.datetime) -> str:
    return f"{t.astimezone(PRAGUE):%d. %m. %Y %H:%M}"


def filled(case_id: str) -> tuple[dict[str, list[Field | None]], dict | None]:
    """The template filled for the meeting (sections) + the AI draft used (or None). 404 for an
    unknown case, 409 before the board's conclusion is confirmed."""
    c = cases.get_case(case_id)
    rec = records.get_record(case_id)
    d = rec.decision
    if not d.confirmed_by:
        raise HTTPException(409, "Závěr zatím není potvrzen")
    sections = parse_template(TEMPLATE.read_text("utf-8"))
    with cases.fhir_server():
        patient = cases.record(cases.meetings()[case_id]).all("Patient")
    fill_header(sections["HEADER"], patient[0] if patient else {}, d, c.board)

    draft = load_draft(case_id)
    ai = draft["fields"] if draft else {}
    question = rec.prep.question or (c.question.text if c.question else NOT_FOUND)
    for f in leaves(sections["BODY"]):
        if f.label == QUESTION:
            f.value = question
        elif f.path in ai:
            f.value, f.sources = ai[f.path]["value"], ai[f.path]["sources"]
        else:
            f.value = NOT_FILLED

    def line(label: str, value: str) -> Field:
        return Field(label, 1, value, False, f"Závěr MDT > {label}")

    response: list[Field | None] = [
        Field("Závěr MDT", 0, "", True, "Závěr MDT"),
        line("Léčebný záměr", INTENTS.get(d.intent or "", "–")),
        line("Zvolený postup", d.plan),
        line("Zdůvodnění", d.rationale),
    ]
    response += [
        line(label, value)
        for label, value in (
            ("Zvažované možnosti", d.options_considered),
            ("Odlišná stanoviska", d.dissent),
            ("Zařazení do klinické studie", d.trial),
            ("Re-evaluace", d.reevaluation and f"{d.reevaluation:%d. %m. %Y}"),
        )
        if value
    ]
    compliant = "Ano" if d.guideline_compliant else f"Ne – {d.deviation_reason}"
    at = f", {_local(d.confirmed_at)}" if d.confirmed_at else ""
    response += [
        line("Soulad s doporučenými postupy", compliant),
        line("Potvrdil(a)", f"{d.confirmed_by}{at} (konsenzus komise)"),
    ]
    sections["RESPONSE"] = response
    return sections, draft


TITLES = {"HEADER": "Multidisciplinární tým a pacient", "BODY": "Podklady pro MDT", "RESPONSE": ""}


def _ai_note(draft: dict | None) -> str:
    if not draft:
        return ""
    created = _local(dt.datetime.fromisoformat(draft["created"]))
    return (
        f"Modře: údaje doplněné AI ({draft['model']}, {created}) ze zdravotního záznamu, "
        "u každého je uveden zdroj. Ke kontrole lékařem; úpravy v DOCX."
    )


def _unsourced(f: Field) -> bool:
    return f.sources is not None and not f.sources and f.value not in (NOT_FOUND, "0")


def to_pdf(title: str, sections: dict, draft: dict | None) -> bytes:
    doc = PdfDoc(f"Závěr zasedání MDT – {title}", footer=f"Závěr zasedání MDT · {title}")
    doc.paragraph([Run("Závěr zasedání MDT", bold=True)], size=16)
    doc.paragraph([Run(title, color=GREY)], size=9)
    if note := _ai_note(draft):
        doc.space(4)
        doc.paragraph([Run(note, color=BLUE)], size=8)
    for name, lines in sections.items():
        doc.rule()
        if TITLES.get(name):
            doc.paragraph([Run(TITLES[name], bold=True)], size=12)
            doc.space(2)
        for f in lines:
            if f is None:
                doc.space(4)
                continue
            indent = 14 * f.depth
            if f.group or not f.value:
                doc.keep(50)
                doc.paragraph([Run(f.label, bold=True)], indent=indent)
                continue
            muted = f.value in (NOT_FOUND, NOT_FILLED)
            color = GREY if muted else BLUE if f.sources is not None else (0, 0, 0)
            doc.paragraph(
                [Run(f"{f.label}: ", bold=True), Run(f.value, color=color)], indent=indent
            )
            if f.sources:
                doc.paragraph(
                    [Run("Zdroj: " + ", ".join(f.sources), color=GREY)],
                    size=6.5,
                    indent=indent,
                    text_offset=10,
                )
            elif _unsourced(f):
                doc.paragraph(
                    [Run(NO_SOURCE, color=(0.75, 0.3, 0))], size=6.5, indent=indent, text_offset=10
                )
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "zaver.pdf"
        doc.save(path)
        return path.read_bytes()


def to_docx(title: str, sections: dict, draft: dict | None) -> bytes:
    doc = Document()
    doc.add_heading("Závěr zasedání MDT", level=1)
    doc.add_paragraph(title).runs[0].font.color.rgb = RGBColor(0x73, 0x73, 0x73)
    if note := _ai_note(draft):
        r = doc.add_paragraph().add_run(note)
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(*(round(c * 255) for c in BLUE))
    for name, lines in sections.items():
        if TITLES.get(name):
            doc.add_heading(TITLES[name], level=2)
        for f in lines:
            if f is None:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14 * f.depth)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f.label + ("" if f.group or not f.value else ": ")).bold = True
            if f.group or not f.value:
                continue
            value = p.add_run(f.value)
            if f.value in (NOT_FOUND, NOT_FILLED):
                value.font.color.rgb = RGBColor(0x73, 0x73, 0x73)
            elif f.sources is not None:
                value.font.color.rgb = RGBColor(*(round(c * 255) for c in BLUE))
            if f.sources or _unsourced(f):
                src = p.add_run(
                    "\nZdroj: " + ", ".join(f.sources) if f.sources else f"\n{NO_SOURCE}"
                )
                src.font.size, src.font.color.rgb = Pt(7), RGBColor(0x73, 0x73, 0x73)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _title(case_id: str) -> str:
    c = cases.get_case(case_id)
    mdt = c.mdt_date and f"MDT {c.mdt_date:%d. %m. %Y}"
    return " · ".join(x for x in (c.label, c.board, mdt) if x)


@router.get("/{case_id}/zaver.pdf")
def zaver_pdf(case_id: str) -> Response:
    sections, draft = filled(case_id)
    return Response(
        to_pdf(_title(case_id), sections, draft),
        media_type="application/pdf",  # inline: the Export step previews it
        headers={"Content-Disposition": f'inline; filename="zaver-mdt-{case_id}.pdf"'},
    )


@router.get("/{case_id}/zaver.docx")
def zaver_docx(case_id: str) -> Response:
    sections, draft = filled(case_id)
    return Response(
        to_docx(_title(case_id), sections, draft),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="zaver-mdt-{case_id}.docx"'},
    )
